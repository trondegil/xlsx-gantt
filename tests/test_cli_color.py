"""
Tests for the ``--color`` CLI option (theme built with
``GanttTheme.from_color``).
"""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from xlsx_gantt import GanttTheme
from xlsx_gantt.cli import main as cli_main

DATE_COL0 = 5  # first date column (column E)

CONFIG = {
    "start_date": "2026-01-12",
    "end_date": "2026-01-18",
    "sections": [{"name": "Build", "tasks": [
        {"name": "Themed bar", "ranges": [{"start": "2026-01-12", "end": "2026-01-13"}]},
        {"name": "Own colour", "ranges": [
            {"start": "2026-01-14", "end": "2026-01-15", "color": "00B050"}]},
    ]}],
}


def _fill(ws, row, col):
    return ws.cell(row=row, column=col).fill.fgColor.rgb[-6:]


@pytest.fixture
def cfg(tmp_path):
    path = tmp_path / "chart.json"
    path.write_text(json.dumps(CONFIG))
    return path


@pytest.mark.parametrize("value", ["2E86C1", "2e86c1", "#2E86C1"])
def test_color_builds_theme_from_base(cfg, tmp_path, value):
    out = tmp_path / "out.xlsx"
    assert cli_main([str(cfg), "-o", str(out), "--color", value]) == 0
    ws = load_workbook(out).active
    expected = GanttTheme.from_color("2E86C1")
    assert _fill(ws, 1, DATE_COL0) == expected.header_bg            # week band
    assert _fill(ws, 4, DATE_COL0) == expected.bar_color            # no range colour
    assert _fill(ws, 5, DATE_COL0 + 2) == "00B050"                  # range colour wins


def test_color_with_week_start(cfg, tmp_path):
    out = tmp_path / "out.xlsx"
    argv = [str(cfg), "-o", str(out), "--color", "2E86C1", "--week-start", "sunday"]
    assert cli_main(argv) == 0
    assert load_workbook(out).active.cell(row=1, column=DATE_COL0).value == "Week 3"


@pytest.mark.parametrize("bad", ["2E86C", "2E86C1F", "GGGGGG", "#12345", "blue"])
def test_invalid_color_rejected(cfg, capsys, bad):
    with pytest.raises(SystemExit):
        cli_main([str(cfg), "--color", bad])
    assert "invalid colour" in capsys.readouterr().err


def test_color_and_theme_are_exclusive(cfg, capsys):
    with pytest.raises(SystemExit):
        cli_main([str(cfg), "--color", "2E86C1", "--theme", "ocean"])
    assert "not allowed with" in capsys.readouterr().err

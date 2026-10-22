"""
Tests for the Total row formulas (``SUM`` of the time estimates and the
overall progress weighted by time estimate) and the Earned / Remaining
rows below it.

The formula text is checked everywhere; the computed values are checked
by letting LibreOffice recalculate the workbook, when it is installed.
"""
from __future__ import annotations

import csv
import io
import shutil
import subprocess
from datetime import date

import pytest
from openpyxl import load_workbook

from xlsx_gantt import GanttChart, GanttStyle

COL_EST, COL_PROG = 3, 4  # columns C and D


def _task(name, est, prog):
    return {"name": name, "time_estimate": est, "progress": prog,
            "ranges": [{"start": date(2026, 1, 12), "end": date(2026, 1, 13)}]}


def _chart(*tasks, **style_kwargs):
    return GanttChart([{"name": "Build", "tasks": list(tasks)}],
                      date(2026, 1, 12), date(2026, 1, 18),
                      style=GanttStyle(**style_kwargs))


def _sheet(chart):
    return load_workbook(io.BytesIO(chart.generate_excel_bytes())).active


def _total_cells(chart):
    ws = _sheet(chart)
    row = next(r for r in range(4, ws.max_row + 1)
               if ws.cell(row=r, column=1).value == "Total")
    return ws.cell(row=row, column=COL_EST).value, ws.cell(row=row, column=COL_PROG).value


def test_formulas_cover_all_task_rows():
    est, prog = _total_cells(_chart(_task("A", 2, 100), _task("B", 8, 0)))
    assert est == "=SUM(C4:C5)"
    assert prog == (
        '=IFERROR(ROUND(SUMPRODUCT(C4:C5,D4:D5)/SUMIF(C4:C5,">0"),1),'
        'IFERROR(ROUND(AVERAGE(D4:D5),1),""))'
    )


def test_no_tasks_has_no_formulas():
    chart = GanttChart([], date(2026, 1, 12), date(2026, 1, 18))
    assert _total_cells(chart) == (0, None)
    assert _sheet(chart).max_row == 4   # no Earned / Remaining rows


def test_earned_and_remaining_rows():
    ws = _sheet(_chart(_task("A", 2, 100), _task("B", 8, 0)))
    # Total is row 6; one blank row, then Earned (8) and Remaining (9).
    assert ws.cell(row=7, column=1).value is None
    assert [[ws.cell(row=r, column=c).value for c in (1, 2, 3)] for r in (8, 9)] == [
        ["Earned hours", "Time estimate × progress, summed",
         "=ROUND(SUMPRODUCT(C4:C5,D4:D5)/100,1)"],
        ["Remaining hours", "Total time estimate − earned", "=C6-C8"],
    ]
    for r in (8, 9):
        for c in (1, 2, 3):
            assert ws.cell(row=r, column=c).font.color.rgb[-6:] == "FFFFFF"


def test_earned_rows_use_style_labels():
    ws = _sheet(_chart(_task("A", 2, 100), earned_label="Earned days",
                       earned_note="Done", remaining_label="Remaining days",
                       remaining_note="Left", summary_fg="808080"))
    assert [ws.cell(row=r, column=c).value for r in (7, 8) for c in (1, 2)] == [
        "Earned days", "Done", "Remaining days", "Left"]
    assert ws.cell(row=7, column=3).font.color.rgb[-6:] == "808080"


# ── Computed values (LibreOffice) ────────────────────────────────────

SCENARIOS = {
    # 2 days done + 8 days not started = 20 %, not the plain mean of 50 %.
    # name:          (tasks, total estimate, progress, earned, remaining)
    "weighted":      ([_task("Small", 2, 100), _task("Big", 8, 0)], 10, 20, 2, 8),
    "rounded":       ([_task("A", 1, 100), _task("B", 2, 0)], 3, 33.3, 1, 2),
    "milestone":     ([_task("A", 3, 50), _task("Go live", None, None)], 3, 50, 1.5, 1.5),
    "no_estimate":   ([_task("A", 4, 25), _task("B", None, 100)], 4, 25, 1, 3),
    "zero_estimate": ([_task("A", 4, 25), _task("B", 0, 100)], 4, 25, 1, 3),
    "text_estimate": ([_task("A", 4, 25), _task("B", "3", 100)], 4, 25, 1, 3),
    # Blank or text progress counts as 0 %, as in Earned / Remaining.
    "text_progress": ([_task("A", 4, 25), _task("B", 6, "N/A")], 10, 10, 1, 9),
    "no_progress":   ([_task("A", 4, 25), _task("B", 6, None)], 10, 10, 1, 9),
    "all_blank":     ([_task("A", 3, None)], 3, 0, 0, 3),
    "plain_mean":    ([_task("A", None, 100), _task("B", None, 0)], 0, 50, 0, 0),
    "nothing":       ([_task("A", None, None)], 0, None, 0, 0),
}


@pytest.fixture(scope="module")
def computed(tmp_path_factory):
    """Recalculate every scenario in one LibreOffice run; return Total rows."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    tmp = tmp_path_factory.mktemp("lo")
    files = []
    for name, (tasks, *_) in SCENARIOS.items():
        path = tmp / f"{name}.xlsx"
        _chart(*tasks).generate_excel(str(path))
        files.append(str(path))
    subprocess.run(
        [soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless",
         "--convert-to", "csv", "--outdir", str(tmp), *files],
        check=True, capture_output=True, timeout=120,
    )
    totals = {}
    for name in SCENARIOS:
        with open(tmp / f"{name}.csv", newline="") as fh:
            rows = {r[0]: r for r in csv.reader(fh) if r}
        # The CSV uses the system locale's decimal separator.
        totals[name] = tuple(
            rows[label][col - 1].replace(",", ".")
            for label, col in (("Total", COL_EST), ("Total", COL_PROG),
                               ("Earned hours", COL_EST),
                               ("Remaining hours", COL_EST))
        )
    return totals


@pytest.mark.parametrize("name", SCENARIOS)
def test_computed_totals(computed, name):
    _, est, prog, earned, remaining = SCENARIOS[name]
    got_est, got_prog, got_earned, got_remaining = computed[name]
    assert float(got_est) == est
    if prog is None:
        assert got_prog == ""
    else:
        assert float(got_prog) == prog
    assert float(got_earned) == earned
    assert float(got_remaining) == remaining

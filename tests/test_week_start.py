"""
Tests for ``GanttStyle.week_start`` (first day of the week) and the
matching ``--week-start`` CLI flag.
"""
from __future__ import annotations

import io
import json
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from xlsx_gantt import GanttChart, GanttStyle, GanttTheme
from xlsx_gantt.chart import _week_of
from xlsx_gantt.cli import main as cli_main

DATE_COL0 = 5  # first date column (column E)


def _sections(start, end):
    return [{"name": "Build", "tasks": [
        {"name": "Code", "ranges": [{"start": start, "end": end}]},
    ]}]


def _load(start, end, **style_kwargs):
    chart = GanttChart(_sections(start, end), start, end,
                       style=GanttStyle(**style_kwargs))
    return load_workbook(io.BytesIO(chart.generate_excel_bytes())).active


def _week_bands(ws):
    """Return the row-1 week bands as sorted ``(first_col, last_col)`` pairs."""
    return sorted(
        (r.min_col, r.max_col) for r in ws.merged_cells.ranges
        if r.min_row == r.max_row == 1 and r.min_col >= DATE_COL0
    )


# ── Week numbering ──────────────────────────────────────────────────────

def test_monday_start_is_iso():
    for d in (date(2026, 1, 1), date(2026, 12, 31), date(2027, 1, 3)):
        assert _week_of(d, 0) == d.isocalendar()[:2]


@pytest.mark.parametrize("d, expected", [
    (date(2025, 12, 28), (2026, 1)),    # Sunday before 1 Jan 2026 (a Thursday)
    (date(2026, 1, 3),   (2026, 1)),    # Saturday ending week 1
    (date(2026, 1, 4),   (2026, 2)),
    (date(2026, 1, 11),  (2026, 3)),
    (date(2026, 12, 26), (2026, 52)),
    (date(2026, 12, 27), (2027, 1)),    # week holding 1 Jan 2027 → 2027 wk 1
    (date(2027, 1, 2),   (2027, 1)),
])
def test_sunday_start_week_numbers(d, expected):
    assert _week_of(d, 6) == expected


def test_week_of_accepts_datetime():
    assert _week_of(datetime(2026, 1, 11, 15, 30), 6) == (2026, 3)


# ── Rendered week bands ─────────────────────────────────────────────────

def test_default_bands_run_monday_to_sunday():
    # Sun 11 Jan – Sat 24 Jan 2026: Monday weeks split off the first Sunday
    ws = _load(datetime(2026, 1, 11), datetime(2026, 1, 24))
    assert _week_bands(ws) == [(6, 12), (13, 18)]
    assert ws.cell(row=1, column=DATE_COL0).value == "Week 2"


def test_sunday_start_bands():
    ws = _load(datetime(2026, 1, 11), datetime(2026, 1, 24), week_start=6)
    assert _week_bands(ws) == [(5, 11), (12, 18)]
    assert ws.cell(row=1, column=5).value == "Week 3"
    assert ws.cell(row=1, column=12).value == "Week 4"


def test_sunday_start_band_spans_new_year():
    ws = _load(datetime(2026, 12, 20), datetime(2027, 1, 2),
               week_start=6, week_label_format="W{week} {year}")
    assert _week_bands(ws) == [(5, 11), (12, 18)]
    assert ws.cell(row=1, column=5).value == "W52 2026"
    assert ws.cell(row=1, column=12).value == "W1 2027"


def test_saturday_start_bands():
    ws = _load(datetime(2026, 1, 10), datetime(2026, 1, 23), week_start=5)
    assert _week_bands(ws) == [(5, 11), (12, 18)]


def test_day_names_unaffected_by_week_start():
    ws = _load(datetime(2026, 1, 11), datetime(2026, 1, 17), week_start=6)
    assert ws.cell(row=2, column=DATE_COL0).value == "Sun"
    assert ws.cell(row=2, column=DATE_COL0 + 1).value == "Mon"


def test_week_start_works_with_themes():
    style = GanttTheme.get("ocean")
    style.week_start = 6
    chart = GanttChart(_sections(datetime(2026, 1, 11), datetime(2026, 1, 17)),
                       datetime(2026, 1, 11), datetime(2026, 1, 17), style=style)
    ws = load_workbook(io.BytesIO(chart.generate_excel_bytes())).active
    assert _week_bands(ws) == [(5, 11)]


# ── Validation ──────────────────────────────────────────────────────────

@pytest.mark.parametrize("bad", [-1, 7, True, "sunday", 6.0, None])
def test_invalid_week_start_rejected(bad):
    with pytest.raises(ValueError, match="week_start"):
        GanttStyle(week_start=bad)


# ── CLI ─────────────────────────────────────────────────────────────────

CLI_CONFIG = {
    "start_date": "2026-01-11",
    "end_date": "2026-01-24",
    "sections": [{"name": "Build", "tasks": [{"name": "Code"}]}],
}


@pytest.mark.parametrize("theme_args", [[], ["--theme", "ocean"]])
def test_cli_week_start_sunday(tmp_path, theme_args):
    cfg = tmp_path / "chart.json"
    out = tmp_path / "out.xlsx"
    cfg.write_text(json.dumps(CLI_CONFIG))
    argv = [str(cfg), "-o", str(out), "--week-start", "sunday", *theme_args]
    assert cli_main(argv) == 0
    assert _week_bands(load_workbook(out).active) == [(5, 11), (12, 18)]


def test_cli_rejects_unknown_week_start(tmp_path):
    cfg = tmp_path / "chart.json"
    cfg.write_text(json.dumps(CLI_CONFIG))
    with pytest.raises(SystemExit):
        cli_main([str(cfg), "--week-start", "sun"])
